#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build the WeChat Channels short-drama metrics workbook.

Reads a JSON detail file and writes an .xlsx with two sheets:
  - 前N集汇总 : one row per 视频号+剧集 with summed metrics
  - 前N集明细 : one row per episode

Input JSON shape:
{
  "rows": [
    {"video_account": "奇镜剧场", "video_id": "sph...", "drama": "剧名",
     "drama_id": "wx...", "episode": 1,
     "play": 195, "like": 0, "comment": 0, "thumb": 1},
    ...
  ]
}

Usage:
  python build_drama_excel.py --detail detail.json --out result.xlsx [--time "2026-09-06 12:00:00"]
"""
import argparse
import datetime
import json
import sys

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is required: pip install openpyxl")


HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")


def to_int(value, field):
    try:
        return int(value)
    except (TypeError, ValueError):
        return 0


def main():
    ap = argparse.ArgumentParser(description="Build short-drama metrics workbook")
    ap.add_argument("--detail", required=True, help="path to JSON detail file")
    ap.add_argument("--out", required=True, help="output .xlsx path")
    ap.add_argument("--time", default=None, help="capture time, default now")
    args = ap.parse_args()

    with open(args.detail, encoding="utf-8") as fh:
        data = json.load(fh)
    rows = data.get("rows") or []
    if not isinstance(rows, list) or not rows:
        sys.exit("No rows found in detail JSON")

    for r in rows:
        r["episode"] = to_int(r.get("episode"), "episode")
        r["play"] = to_int(r.get("play"), "play")
        r["like"] = to_int(r.get("like"), "like")
        r["comment"] = to_int(r.get("comment"), "comment")
        r["thumb"] = to_int(r.get("thumb"), "thumb")

    # sort: by drama fields, then episode number
    rows.sort(key=lambda r: (
        str(r.get("video_account", "")),
        str(r.get("video_id", "")),
        str(r.get("drama", "")),
        r["episode"],
    ))

    # group per drama
    dramas = {}
    for r in rows:
        key = (str(r.get("video_account", "")), str(r.get("video_id", "")),
               str(r.get("drama", "")), str(r.get("drama_id", "")))
        dramas.setdefault(key, []).append(r)

    capture_time = args.time or datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    wb = openpyxl.Workbook()

    # --- summary sheet ---
    ws = wb.active
    ws.title = "前N集汇总"
    sum_headers = ["视频号", "视频号ID", "剧集名称", "剧集ID", "收录集数",
                   "播放量合计", "喜欢量合计", "评论量合计", "点赞量合计", "抓取时间"]
    ws.append(sum_headers)
    for key, ep_rows in dramas.items():
        account, vid, drama, did = key
        nums = [e["episode"] for e in ep_rows]
        ep_label = f"第{min(nums)}-{max(nums)}集" if len(nums) > 1 else f"第{nums[0]}集"
        ws.append([
            account, vid, drama, did, ep_label,
            sum(e["play"] for e in ep_rows),
            sum(e["like"] for e in ep_rows),
            sum(e["comment"] for e in ep_rows),
            sum(e["thumb"] for e in ep_rows),
            capture_time,
        ])

    # --- detail sheet ---
    ws2 = wb.create_sheet("前N集明细")
    det_headers = ["视频号", "视频号ID", "剧集名称", "剧集ID", "集数",
                   "播放量", "喜欢量", "评论量", "点赞量"]
    ws2.append(det_headers)
    for r in rows:
        ws2.append([
            r.get("video_account", ""), r.get("video_id", ""), r.get("drama", ""),
            r.get("drama_id", ""), f"第{r['episode']}集",
            r["play"], r["like"], r["comment"], r["thumb"],
        ])

    # style both sheets
    for sheet, widths in ((ws, [16, 16, 30, 30, 12, 14, 14, 14, 14, 20]),
                          (ws2, [16, 16, 30, 30, 10, 12, 12, 12, 12])):
        for c in sheet[1]:
            c.fill = HEADER_FILL
            c.font = HEADER_FONT
            c.border = BORDER
            c.alignment = CENTER
        for row in sheet.iter_rows(min_row=2):
            for c in row:
                c.border = BORDER
                c.alignment = CENTER
        for i, w in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(i)].width = w

    wb.save(args.out)
    print(f"saved: {args.out}")
    print(f"dramas={len(dramas)} episode_rows={len(rows)}")


if __name__ == "__main__":
    main()